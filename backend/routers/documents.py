import logging
import os
import re
import uuid

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy import func
from sqlalchemy.orm import Session
from pypdf import PdfReader
from langchain_text_splitters import RecursiveCharacterTextSplitter

from backend.database import get_db
from backend.deps import get_current_user
from backend.models.user import User
from backend.models.document import Document
from backend.models.chunk import Chunk
from backend.services import activity
from backend.services.activity import log_activity
from backend.services.rag import AIServiceError, embed_chunks

router = APIRouter(prefix="/documents", tags=["documents"])

logger = logging.getLogger(__name__)

UPLOAD_DIR = os.getenv("UPLOAD_DIR") or os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads"
)
os.makedirs(UPLOAD_DIR, exist_ok=True)

ALLOWED_FILE_TYPES = {"pdf", "txt", "md", "docx", "pptx"}
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "20"))
READ_CHUNK_BYTES = 1024 * 1024


class UnreadableDocumentError(Exception):
    """The file could not be parsed or contained no usable text."""


def _extract_text(file_path: str, file_type: str) -> str:
    try:
        if file_type == "pdf":
            reader = PdfReader(file_path)
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if file_type in ("txt", "md"):
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        if file_type == "docx":
            import docx
            doc = docx.Document(file_path)
            return "\n".join(para.text for para in doc.paragraphs)
        if file_type == "pptx":
            import pptx
            presentation = pptx.Presentation(file_path)
            return "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if hasattr(shape, "text")
            )
    except Exception as e:
        raise UnreadableDocumentError(f"Could not parse {file_type} file: {e}") from e
    raise UnreadableDocumentError(f"Unsupported file type: {file_type}")


def _remove_file(file_path: str):
    try:
        os.remove(file_path)
    except FileNotFoundError:
        pass
    except OSError:
        logger.exception("Could not remove file %s", file_path)


def _cleanup_failed_upload(db: Session, document_id: int | None, file_path: str):
    """Best-effort removal of everything a failed upload may have created.

    Each step is guarded on its own so a cleanup failure is logged but never
    replaces the original error.
    """
    try:
        db.rollback()
    except Exception:
        logger.exception("Rollback failed while cleaning up upload")

    if document_id is not None:
        # Deleting the document row cascades to its chunks and their vectors
        try:
            doc = db.query(Document).filter(Document.id == document_id).first()
            if doc:
                db.delete(doc)
                db.commit()
        except Exception:
            db.rollback()
            logger.exception("Could not remove document row %s", document_id)

    _remove_file(file_path)


@router.post("/upload")
async def upload_document(
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    # Keep only the base name of the client-supplied filename (strip any path parts)
    original_name = os.path.basename((file.filename or "").replace("\\", "/")) or "upload"
    file_ext = os.path.splitext(original_name)[1].lower()
    if not re.fullmatch(r"\.[a-z0-9]{1,10}", file_ext):
        file_ext = ""
    file_type = file_ext[1:]

    if file_type not in ALLOWED_FILE_TYPES:
        raise HTTPException(
            status_code=415,
            detail=f"Unsupported file type. Allowed: {', '.join(sorted(ALLOWED_FILE_TYPES))}",
        )

    # Save under a unique generated name so uploads can't escape UPLOAD_DIR
    # or overwrite another user's file with the same name. Stream it to disk
    # and stop as soon as it passes the size limit.
    file_path = os.path.join(UPLOAD_DIR, f"{uuid.uuid4().hex}{file_ext}")
    max_bytes = MAX_UPLOAD_MB * 1024 * 1024
    size = 0
    too_large = False
    try:
        with open(file_path, "wb") as f:
            while piece := await file.read(READ_CHUNK_BYTES):
                size += len(piece)
                if size > max_bytes:
                    too_large = True
                    break
                f.write(piece)
    except Exception:
        _remove_file(file_path)
        raise
    if too_large:
        _remove_file(file_path)
        raise HTTPException(status_code=413, detail=f"File is larger than {MAX_UPLOAD_MB} MB")

    document_id = None
    try:
        # Create document record
        db_doc = Document(
            user_id=user.id,
            filename=original_name,
            file_type=file_type,
            upload_path=file_path
        )
        db.add(db_doc)
        db.commit()
        db.refresh(db_doc)
        document_id = db_doc.id

        # Extract and chunk text
        text = _extract_text(file_path, file_type)
        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
        )
        texts = [t for t in text_splitter.split_text(text) if t.strip()]
        if not texts:
            raise UnreadableDocumentError("No text could be extracted")

        # Embed the chunks, then save them with their vectors
        db_chunks = [
            Chunk(
                document_id=db_doc.id,
                user_id=user.id,
                content=chunk_text,
                chunk_index=idx
            )
            for idx, chunk_text in enumerate(texts)
        ]
        embed_chunks(db_chunks)
        db.add_all(db_chunks)

        log_activity(db, user.id, activity.DOCUMENT_UPLOADED, db_doc.id, {
            "filename": original_name,
            "chunks": len(db_chunks),
        })
        db.commit()

    except UnreadableDocumentError as e:
        logger.info("Rejected upload %r: %s", original_name, e)
        _cleanup_failed_upload(db, document_id, file_path)
        raise HTTPException(status_code=422, detail="No readable text could be extracted from this file")
    except AIServiceError:
        _cleanup_failed_upload(db, document_id, file_path)
        raise  # turned into a generic 502 by the app's exception handler
    except Exception:
        logger.exception("Unexpected error while processing upload %r", original_name)
        _cleanup_failed_upload(db, document_id, file_path)
        raise HTTPException(status_code=500, detail="Failed to process the uploaded document")

    return {"message": "Document uploaded and processed successfully", "document_id": db_doc.id, "chunks_count": len(db_chunks)}

@router.get("/")
def get_documents(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    docs = (
        db.query(Document)
        .filter(Document.user_id == user.id)
        .order_by(Document.created_at.desc(), Document.id.desc())
        .all()
    )
    chunk_counts = dict(
        db.query(Chunk.document_id, func.count(Chunk.id))
        .filter(Chunk.user_id == user.id)
        .group_by(Chunk.document_id)
        .all()
    )
    return [
        {
            "id": d.id,
            "filename": d.filename,
            "file_type": d.file_type,
            "chunk_count": chunk_counts.get(d.id, 0),
            "created_at": d.created_at,
        }
        for d in docs
    ]

@router.delete("/{document_id}")
def delete_document(document_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == document_id, Document.user_id == user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    # Remove file from disk
    _remove_file(doc.upload_path)

    # Delete from the database (cascade handles chunks and their vectors, flashcard sets and fill-in-the-blank sets)
    db.delete(doc)
    db.commit()

    return {"message": "Document deleted successfully"}
